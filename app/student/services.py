"""Student services — student auth, CRUD, and XLSX import."""
from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.auth.models import User
from app.auth.repositories import UserRepository
from app.infra.audit import log as audit_log
from app.infra.passwords import hash_password, is_legacy_hash, verify_password
from app.student.dto import (
    StudentCreateIn,
    StudentLoginIn,
    StudentRegisterIn,
    StudentUpdateIn,
)
from app.student.models import Student
from app.student.repositories import StudentRepository
from framework.base import ServiceBase
from framework.exceptions import Conflict, DomainException, NotFound, ValidationFailed

REQUIRED_COLUMNS = ("first_name", "last_name", "email", "username", "password", "category")
OPTIONAL_COLUMNS = ("school", "city", "country", "status")
_ACTIVE_VALUES = {"aktif", "active", "1", "true", "yes"}


def branding() -> dict[str, Any]:
    """Public branding payload for the student login screen."""
    return {
        "title": "Tantangan Bebras",
        "subtitle": "Student Login",
        "description": (
            "Sign in to take the Bebras Computational Thinking Challenge "
            "according to your category (Cub / Scout / Rover)."
        ),
        "registerUrl": "/student-register",
        "passwordMinLength": 6,
        "categoryOptions": ["cub", "scout", "rover"],
    }


class StudentAuthService(ServiceBase[StudentRepository]):
    """Business logic for student self-registration and login."""

    repository_class = StudentRepository

    def __init__(
        self,
        repo: StudentRepository | None = None,
        users: UserRepository | None = None,
    ):
        super().__init__(repo)
        self._users = users or UserRepository()

    def register(self, dto: StudentRegisterIn) -> tuple[User, Student]:
        try:
            if self._users.find_by_username(dto.username):
                raise Conflict("username_already_registered")
            if self._users.find_by_email(dto.email):
                raise Conflict("email_already_registered")
        except DomainException as exc:
            audit_log(
                "student_register_failure",
                resource="student-auth/register",
                details={"username": dto.username, "reason": exc.message},
            )
            self._commit()
            raise

        user = User(
            username=dto.username,
            email=dto.email,
            first_name=dto.first_name,
            last_name=dto.last_name,
            full_name=f"{dto.first_name} {dto.last_name}",
            password=hash_password(dto.password),
            active=True,
            app_role="user",
        )
        self._users.add(user)
        self._flush()

        student = Student(
            user_id=user.id,
            category=dto.category,
            school=dto.school,
            city=dto.city,
            country=dto.country,
            status="Inactive",
        )
        self._repo.add(student)
        audit_log(
            "student_register_success",
            user_id=user.id,
            resource="student-auth/register",
            details={"username": user.username, "category": dto.category},
        )
        self._commit()
        return user, student

    def authenticate(self, dto: StudentLoginIn) -> tuple[User, Student | None] | None:
        user = self._users.find_by_username(dto.username)
        student = self._repo.find_by_user_id(user.id) if user else None
        valid = (
            user is not None
            and user.active
            and verify_password(user.password, dto.password)
            and not (student is not None and student.disabled)
        )
        if not valid:
            audit_log(
                "student_login_failure",
                resource="student-auth/login",
                details={"username": dto.username, "reason": "invalid_credentials"},
            )
            self._commit()
            return None
        if is_legacy_hash(user.password):
            user.password = hash_password(dto.password)
        audit_log(
            "student_login_success",
            user_id=user.id,
            resource="student-auth/login",
            details={"username": user.username},
        )
        self._commit()
        return user, student

    def record_logout(self, user_id: str) -> None:
        audit_log("student_logout", user_id=user_id, resource="student-auth/logout")
        self._commit()

    def get_me(self, user_id: str) -> tuple[User, Student | None]:
        user = self._users.get(user_id)
        if user is None:
            raise NotFound("user_not_found")
        return user, self._repo.find_by_user_id(user_id)


class StudentService(ServiceBase[StudentRepository]):
    """Business logic for admin student management and import."""

    repository_class = StudentRepository

    def __init__(
        self,
        repo: StudentRepository | None = None,
        users: UserRepository | None = None,
    ):
        super().__init__(repo)
        self._users = users or UserRepository()

    def list_students(self, **filters) -> tuple[list[Student], int]:
        return self._repo.list_students(**filters)

    def get_student(self, user_id: str) -> tuple[User, Student]:
        student = self._repo.find_by_user_id(user_id)
        if student is None:
            raise NotFound("student_not_found")
        return student.user, student

    def create_student(self, dto: StudentCreateIn, *, actor_id: str | None) -> Student:
        if self._users.find_by_username(dto.username):
            raise Conflict("username_already_registered")
        if self._users.find_by_email(dto.email):
            raise Conflict("email_already_registered")

        user = self._build_user(dto)
        self._users.add(user)
        self._flush()

        student = Student(
            user_id=user.id,
            category=dto.category,
            school=dto.school,
            city=dto.city,
            country=dto.country,
            status=dto.status,
        )
        self._repo.add(student)
        audit_log(
            "student_created",
            user_id=actor_id,
            resource=f"student/{user.id}",
            details={"category": dto.category, "status": dto.status},
        )
        self._commit()
        return student

    def update_student(self, user_id: str, dto: StudentUpdateIn, *, actor_id: str) -> tuple[User, Student]:
        user, student = self.get_student(user_id)
        data = dto.model_dump(exclude_unset=True)

        if "username" in data and data["username"] != user.username:
            existing = self._users.find_by_username(data["username"])
            if existing is not None and existing.id != user.id:
                raise Conflict("username_already_registered")
            user.username = data["username"]
        if "email" in data and data["email"] != user.email:
            existing = self._users.find_by_email(data["email"])
            if existing is not None and existing.id != user.id:
                raise Conflict("email_already_registered")
            user.email = data["email"]
        if "first_name" in data:
            user.first_name = data["first_name"]
            user.full_name = f"{user.first_name} {user.last_name}".strip()
        if "last_name" in data:
            user.last_name = data["last_name"]
            user.full_name = f"{user.first_name} {user.last_name}".strip()
        if data.get("password"):
            user.password = hash_password(data["password"])
        for field in ("category", "school", "city", "country", "status"):
            if field in data:
                setattr(student, field, data[field])

        audit_log(
            "student_updated",
            user_id=actor_id,
            resource=f"student/{user_id}",
            details={"changed": sorted(data)},
        )
        self._commit()
        return user, student

    def soft_delete(self, user_id: str, *, actor_id: str) -> None:
        user, student = self.get_student(user_id)
        student.disabled = True
        student.disabled_at = datetime.now(UTC)
        user.active = False
        audit_log("student_soft_deleted", user_id=actor_id, resource=f"student/{user_id}")
        self._commit()

    def restore(self, user_id: str, *, actor_id: str) -> tuple[User, Student]:
        user, student = self.get_student(user_id)
        student.disabled = False
        student.disabled_at = None
        user.active = True
        audit_log("student_restored", user_id=actor_id, resource=f"student/{user_id}")
        self._commit()
        return user, student

    def activate(self, user_id: str, *, actor_id: str) -> tuple[User, Student]:
        user, student = self.get_student(user_id)
        if student.disabled:
            raise NotFound("student_disabled")
        student.status = "Active"
        audit_log("student_activated", user_id=actor_id, resource=f"student/{user_id}")
        self._commit()
        return user, student

    def purge(self, user_id: str, *, actor_id: str) -> None:
        user, student = self.get_student(user_id)
        self._repo.delete(student)
        self._users.delete(user)
        audit_log("student_purged", user_id=actor_id, resource=f"student/{user_id}")
        self._commit()

    # --- import -------------------------------------------------------
    def parse_workbook(self, stream) -> tuple[list[dict], dict | None]:
        """Parse an XLSX upload into row dicts. Returns (rows, file_error)."""
        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
        except Exception as exc:
            return [], {"row": 0, "message": f"xlsx_unreadable: {exc}"}

        try:
            sheet = workbook.active
            if sheet is None:
                return [], {"row": 0, "message": "xlsx_empty_sheet"}

            rows_iter = sheet.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if not header or all(cell is None for cell in header):
                return [], {"row": 0, "message": "xlsx_missing_header"}

            columns = {
                str(cell).strip().lower().replace(" ", "_"): idx
                for idx, cell in enumerate(header)
                if cell is not None and str(cell).strip()
            }
            missing = [c for c in REQUIRED_COLUMNS if c not in columns]
            if missing:
                return [], {"row": 0, "message": f"missing_required_columns: {','.join(missing)}"}

            rows: list[dict] = []
            for excel_row in rows_iter:
                if all(cell is None or str(cell).strip() == "" for cell in excel_row):
                    continue
                record = {}
                for key, idx in columns.items():
                    value = excel_row[idx] if idx < len(excel_row) else None
                    record[key] = "" if value is None else str(value).strip()
                rows.append(record)

            if not rows:
                return [], {"row": 0, "message": "no_data_rows"}
            return rows, None
        finally:
            workbook.close()

    def import_rows(self, rows: list[dict], *, actor_id: str | None) -> dict[str, Any]:
        inserted = 0
        updated = 0
        errors: list[dict] = []

        for offset, record in enumerate(rows):
            excel_row = offset + 2
            payload = dict(record)
            payload["status"] = _resolve_status(record.get("status"))
            try:
                dto = StudentCreateIn.validate(payload)
            except ValidationFailed as exc:
                errors.append({"row": excel_row, "message": f"validation: {exc.message}"})
                continue

            existing_by_email = self._users.find_by_email(dto.email)
            if existing_by_email is not None and existing_by_email.username != dto.username:
                errors.append({"row": excel_row, "message": "email_already_registered"})
                continue

            existing_by_username = self._users.find_by_username(dto.username)
            if existing_by_username is not None and self._repo.find_by_user_id(
                existing_by_username.id
            ) is None:
                errors.append({"row": excel_row, "message": "username_already_registered"})
                continue
            try:
                if existing_by_username is not None:
                    self._import_update(existing_by_username, dto, actor_id=actor_id)
                    updated += 1
                else:
                    self.create_student(dto, actor_id=actor_id)
                    inserted += 1
            except DomainException as exc:
                errors.append({"row": excel_row, "message": exc.message})

        audit_log(
            "student_import",
            user_id=actor_id,
            resource="student/import",
            details={"inserted": inserted, "updated": updated, "errors": len(errors)},
        )
        self._commit()
        return {"inserted": inserted, "updated": updated, "errors": errors}

    def _import_update(self, user: User, dto: StudentCreateIn, *, actor_id: str | None) -> None:
        student = self._repo.find_by_user_id(user.id)
        if student is None:
            student = Student(user_id=user.id, category=dto.category, status="Inactive")
            self._repo.add(student)
        user.first_name = dto.first_name
        user.last_name = dto.last_name
        user.full_name = f"{dto.first_name} {dto.last_name}"
        user.email = dto.email
        student.category = dto.category
        student.school = dto.school
        student.city = dto.city
        student.country = dto.country
        student.status = dto.status
        audit_log(
            "student_updated",
            user_id=actor_id,
            resource=f"student/{user.id}",
            details={"changed": ["import"]},
        )
        self._commit()

    @staticmethod
    def _build_user(dto: StudentCreateIn | StudentRegisterIn) -> User:
        return User(
            username=dto.username,
            email=dto.email,
            first_name=dto.first_name,
            last_name=dto.last_name,
            full_name=f"{dto.first_name} {dto.last_name}",
            password=hash_password(dto.password),
            active=True,
            app_role="user",
        )


def build_template() -> bytes:
    """Generate the student import XLSX template."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "student"
    sheet.append(list(REQUIRED_COLUMNS) + list(OPTIONAL_COLUMNS))
    sheet.append(
        ["Budi", "Santoso", "budi@example.test", "budi", "Siswa1234", "scout",
         "SMPN 1 Bandung", "Bandung", "Indonesia", "Inactive"]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _resolve_status(raw: str | None) -> str:
    return "Active" if (raw or "").strip().lower() in _ACTIVE_VALUES else "Inactive"


__all__ = [
    "REQUIRED_COLUMNS",
    "StudentAuthService",
    "StudentService",
    "branding",
    "build_template",
]
